from datetime import date
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from app.services.citations import render_citations


def _paragraphs(text):
    """Split model text into paragraphs."""
    return [
        part.strip()
        for part in str(text or "").split("\n")
        if part.strip()
    ]


def build_docx(result: dict) -> bytes:
    """
    Build a Word document for a completed research result.

    Uses the literature review when available, otherwise
    the synthesis. Citations are rendered in APA style and
    followed by the reference list.
    """

    sources = result.get("sources") or {}

    def cite(text):
        return render_citations(text, sources)

    document = Document()

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    review = result.get("literature_review") or {}
    synthesis = result.get("synthesis") or {}

    title = (
        review.get("title")
        or f"Literature Review: {result.get('question', '')}"
    )

    heading = document.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"Research question: {result.get('question', '')}").italic = True
    meta.add_run(f"\nGenerated {date.today():%d %B %Y}")

    def add_section(name, text):
        paragraphs = _paragraphs(text)

        if not paragraphs:
            return

        document.add_heading(name, level=1)

        for paragraph in paragraphs:
            document.add_paragraph(cite(paragraph))

    if review:
        add_section("Abstract", review.get("abstract"))
        add_section("Introduction", review.get("introduction"))

        for theme in review.get("themes") or []:
            add_section(theme.get("title") or "Theme", theme.get("discussion"))

        gaps = review.get("research_gaps") or []

        if gaps:
            document.add_heading("Research Gaps", level=1)

            for gap in gaps:
                document.add_paragraph(cite(gap), style="List Bullet")

        add_section("Conclusion", review.get("conclusion"))

    else:
        add_section("Summary", synthesis.get("summary"))

        findings = synthesis.get("key_findings") or []

        if findings:
            document.add_heading("Key Findings", level=1)

            for finding in findings:
                document.add_paragraph(cite(finding), style="List Bullet")

        limitations = synthesis.get("limitations") or []

        if limitations:
            document.add_heading("Limitations", level=1)

            for item in limitations:
                document.add_paragraph(cite(item), style="List Bullet")

    references = result.get("references") or []

    if references:
        document.add_heading("References", level=1)

        for entry in references:
            paragraph = document.add_paragraph(entry)

            # APA hanging indent.
            paragraph.paragraph_format.left_indent = Inches(0.5)
            paragraph.paragraph_format.first_line_indent = Inches(-0.5)

    note = document.add_paragraph()
    note.add_run(
        "Generated with Orion Research. Verify every citation "
        "against the original source before submission."
    ).italic = True

    buffer = BytesIO()
    document.save(buffer)

    return buffer.getvalue()
