from datetime import datetime
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from app.lab.lab_store import lab_store


def _clock(timestamp):
    try:
        return datetime.fromisoformat(timestamp).strftime("%H:%M")
    except (TypeError, ValueError):
        return str(timestamp or "")


def build_notebook_docx(session, entries) -> bytes:
    notebook = session.get("notebook") or {}

    document = Document()

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    title = document.add_heading(
        notebook.get("title") or session["title"],
        level=0,
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    started = session.get("created_at") or ""
    ended = session.get("ended_at") or ""

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(
        f"Session: {started.replace('T', ' ')}"
        + (f" to {_clock(ended)}" if ended else "")
    ).italic = True

    def paragraph_section(name, text):
        if text:
            document.add_heading(name, level=1)
            document.add_paragraph(text)

    def list_section(name, items, style="List Bullet"):
        if items:
            document.add_heading(name, level=1)

            for item in items:
                document.add_paragraph(item, style=style)

    paragraph_section("Objective", notebook.get("objective"))

    background = notebook.get("background")

    if background and background.get("text"):
        paragraph_section("Background", background["text"])

    list_section("Materials", notebook.get("materials"))
    list_section("Equipment", notebook.get("equipment"))

    procedure = notebook.get("procedure") or []

    if procedure:
        document.add_heading("Procedure", level=1)

        for item in procedure:
            paragraph = document.add_paragraph(style="List Number")

            if item.get("time"):
                paragraph.add_run(f"{item['time']}  ").bold = True

            paragraph.add_run(item.get("step", ""))

    list_section("Observations", notebook.get("observations"))
    list_section("Deviations from Protocol", notebook.get("deviations"))
    list_section("Safety Notes", notebook.get("safety_notes"))
    paragraph_section("Summary", notebook.get("summary"))

    if background and background.get("references"):
        document.add_heading("References", level=1)

        for entry in background["references"]:
            paragraph = document.add_paragraph(entry)
            paragraph.paragraph_format.left_indent = Inches(0.5)
            paragraph.paragraph_format.first_line_indent = Inches(-0.5)

    photographed = [entry for entry in entries if entry["has_image"]]

    if photographed:
        document.add_page_break()
        document.add_heading("Photographic Record", level=1)

        for entry in photographed:
            path = lab_store.image_path(session["session_id"], entry["entry_id"])

            if not path.exists():
                continue

            document.add_picture(str(path), width=Inches(4))

            caption = document.add_paragraph()
            caption.add_run(f"{_clock(entry['started_at'])}  ").bold = True
            caption.add_run(entry["activity"]).italic = True

    note = document.add_paragraph()
    note.add_run(
        "Generated with Orion Lab Copilot from an automatic "
        "camera log. Review and sign before filing."
    ).italic = True

    buffer = BytesIO()
    document.save(buffer)

    return buffer.getvalue()
